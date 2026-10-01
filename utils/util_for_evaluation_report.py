"""Export raw and legacy-normalized metrics without requiring pandas."""

from __future__ import annotations

import csv
import os
import re
import tempfile
from collections import defaultdict
from copy import deepcopy
from pathlib import Path

from .util import write_json
from .util_for_evaluation_metrics import winrate_normalize_one_student


_COUNT_PREFIXES = ("annotation_error_count_", "simulation_stop_reason_count_")
_DIAGNOSTIC_PREFIXES = {
    "diagnostics": "diagnostics_",
    "annotation_error_counts": _COUNT_PREFIXES[0],
    "simulation_stop_reasons": _COUNT_PREFIXES[1],
}


def normalize_results(results: list[dict]) -> list[dict]:
    """Preserve raw ARG, including when re-normalizing a saved report."""
    rows = deepcopy(results)
    groups = defaultdict(list)
    for row in rows:
        row.setdefault("raw_avg_doc_answer_reachability_gain", row["avg_doc_answer_reachability_gain"])
        groups[(row["target_data"], row["student_model_name"])].append(row)
    for group in groups.values():
        teachers = [row["teacher_model_name"] for row in group]
        values = [row["raw_avg_doc_answer_reachability_gain"] for row in group]
        if len(set(teachers)) != len(teachers):
            raise ValueError("Teacher IDs must be unique")
        unavailable = sum(value is None for value in values)
        for row in group:
            row.setdefault("diagnostics", {})["arg_comparison_unavailable_teacher_count"] = unavailable
        if unavailable:
            for row in group:
                row["avg_doc_answer_reachability_gain"] = None
            continue
        scores = winrate_normalize_one_student(teachers, values)
        for row in group:
            row["avg_doc_answer_reachability_gain"] = scores[row["teacher_model_name"]]
    return rows


def _flat_row(row: dict) -> dict:
    """Keep validity/error counts beside the scalar scores in CSV and Excel."""
    flattened = {}
    for key, value in row.items():
        if key.startswith("feedback_count2"):
            for count, metric in sorted(value.items(), key=lambda item: int(item[0])):
                flattened[key.replace("2", f"_{count}_", 1)] = metric
        elif key in _DIAGNOSTIC_PREFIXES:
            for name, metric in sorted(value.items()):
                if isinstance(metric, (dict, list)):
                    raise ValueError(f"{key}.{name} must be a scalar report value")
                flattened[f"{_DIAGNOSTIC_PREFIXES[key]}{name}"] = metric
            if key == "annotation_error_counts":
                flattened["annotation_error_total_count"] = sum(value.values())
            elif key == "simulation_stop_reasons":
                flattened["simulation_stop_total_count"] = sum(value.values())
        elif not isinstance(value, (dict, list)):
            flattened[key] = value
    return flattened


def _atomic_export(path: Path, write) -> None:
    descriptor, name = tempfile.mkstemp(prefix=f".{path.stem}.", suffix=path.suffix, dir=path.parent)
    os.close(descriptor)
    temporary = Path(name)
    try:
        write(temporary)
        os.replace(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)


def write_reports(path: Path, rows: list[dict], *, write_excel: bool = True) -> None:
    """Write full JSON and CSV/Excel with scalar metrics and diagnostic counts.

    Missing categories in count maps mean zero. An explicitly unavailable
    scalar remains blank; per-problem feedback counts stay in separate sheets.
    """
    flat = [_flat_row(row) for row in rows]
    columns = list(dict.fromkeys(key for row in flat for key in row))
    count_columns = [column for column in columns if column.startswith(_COUNT_PREFIXES)]
    for row in flat:
        for column in count_columns:
            row.setdefault(column, 0)
    write_json({"schema_version": 1, "results": rows}, path / "metrics.json")

    def write_csv(destination):
        with destination.open("w", newline="", encoding="utf-8-sig") as stream:
            writer = csv.DictWriter(stream, fieldnames=columns)
            writer.writeheader()
            writer.writerows(flat)

    _atomic_export(path / "metrics.csv", write_csv)
    if not write_excel:
        return
    try:
        from openpyxl import Workbook
    except ImportError as exc:
        raise RuntimeError(
            "Excel export requires openpyxl; install requirements-evaluation.txt "
            "or set evaluation.write_excel=false. JSON and CSV were saved."
        ) from exc

    workbook = Workbook()
    summary = workbook.active
    summary.title = "all_results"

    def append(sheet, values):
        sheet.append(values)
        # Dataset/model names are data, not spreadsheet formulae.
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"

    append(summary, columns)
    for row in flat:
        append(summary, [row.get(column) for column in columns])
    groups = defaultdict(list)
    for row, item in zip(rows, flat):
        groups[(row["target_data"], row["teacher_model_name"])].append((row, item))
    used = {"all_results"}
    for index, ((dataset, teacher), group) in enumerate(sorted(groups.items()), start=1):
        suffix = f"_{index}"
        title = re.sub(r"[\\/*?:\[\]]", "_", f"{dataset}_{teacher}")[:31 - len(suffix)] + suffix
        if title in used:
            raise ValueError("Duplicate workbook sheet title")
        used.add(title)
        sheet = workbook.create_sheet(title)
        append(sheet, columns)
        for row, item in group:
            append(sheet, [item.get(column) for column in columns])
        detail = workbook.create_sheet(f"problems_{index}")
        append(detail, ["q_id", *[row["student_model_name"] for row, _ in group]])
        feedback_maps = [row["q_id2feedback_count"] for row, _ in group]
        q_ids = sorted({int(q_id) for feedback in feedback_maps for q_id in feedback})
        for q_id in q_ids:
            append(detail, [q_id, *[
                feedback.get(q_id, feedback.get(str(q_id))) for feedback in feedback_maps
            ]])
    try:
        _atomic_export(path / "metrics.xlsx", workbook.save)
    finally:
        workbook.close()
